from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUTPUT = r"E:\Suilli10086\web_build\docs\页面布局课程教学网站项目说明.docx"

NAVY = "1F4E79"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "F4F8FC"
LIGHT_GRAY = "D9D9D9"
TEXT = "000000"


def set_run_font(run, name="SimSun", size=10.5, bold=False, color=TEXT):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=LIGHT_GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, "SimSun", 9)


def set_paragraph_format(paragraph, before=0, after=0, line=18, first_indent=True):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = Pt(line)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    if first_indent:
        pf.first_line_indent = Cm(0.74)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def add_body(doc, text, before=0, after=7, bold_prefix=None):
    p = doc.add_paragraph()
    set_paragraph_format(p, before, after, 19, True)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, "SimSun", 10.5, True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r, "SimSun", 10.5)
    else:
        r = p.add_run(text)
        set_run_font(r, "SimSun", 10.5)
    return p


def add_heading(doc, text, level=1):
    style = "Heading 1" if level == 1 else "Heading 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    if level == 1:
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(9)
        p.paragraph_format.line_spacing = Pt(24)
    else:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = Pt(20)
    r = p.add_run(text)
    set_run_font(r, "Microsoft YaHei", 15 if level == 1 else 12, True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = Pt(18)
    r = p.add_run(text)
    set_run_font(r, "SimSun", 10.5)
    return p


def add_table(doc, headers, rows, widths, font_size=9.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, label in enumerate(headers):
        cell = header.cells[i]
        set_cell_width(cell, widths[i])
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(15)
        r = p.add_run(label)
        set_run_font(r, "Microsoft YaHei", font_size, True, "FFFFFF")
    for ri, row_values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            cell = cells[i]
            set_cell_width(cell, widths[i])
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(15)
            r = p.add_run(str(value))
            set_run_font(r, "SimSun", font_size)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    return table


def configure_document(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.45)
    sec.bottom_margin = Cm(2.3)
    sec.left_margin = Cm(2.55)
    sec.right_margin = Cm(2.55)
    sec.header_distance = Cm(1.25)
    sec.footer_distance = Cm(1.25)

    normal = doc.styles["Normal"]
    normal.font.name = "SimSun"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = Pt(19)
    normal.paragraph_format.space_after = Pt(7)

    for name in ("Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True

    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("第 ")
    set_run_font(r, "SimSun", 9)
    add_page_field(p)
    r = p.add_run(" 页")
    set_run_font(r, "SimSun", 9)

    props = doc.core_properties
    props.title = "页面布局课程教学网站设计与实现"
    props.subject = "可视化课程作业项目说明"
    props.author = ""
    props.keywords = "页面布局 教学网站 可视化 双语 交互"


def main():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(16)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.line_spacing = Pt(34)
    r = title.add_run("页面布局课程教学网站设计与实现")
    set_run_font(r, "Microsoft YaHei", 22, True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(22)
    subtitle.paragraph_format.line_spacing = Pt(21)
    r = subtitle.add_run("可视化课程作业项目说明")
    set_run_font(r, "Microsoft YaHei", 13)

    add_heading(doc, "摘要", 1)
    add_body(doc, "本项目完成了一个名为《页面布局之道》的双语交互式教学网站，面向学习 CSS 布局、响应式设计和信息层级的初学者。网站将课程要求中的讲解、关键代码和动态演示统一组织到每一个知识小节中，并以 NeoCampus 新生数据看板为贯穿案例，展示同一份语义化 HTML 与同一份数据如何通过逐层增加 CSS 能力，最终形成可交付的页面。课程内容按照内容、间距、尺度、结构、视觉和适应六层展开，避免将布局知识拆成互不关联的属性说明。项目使用 Astro、Svelte、TypeScript、Zod、D3.js 和 p5.js 构建，采用静态生成与交互孤岛结合的架构，同时通过路由级国际化、内容校验、设计令牌检查和单元测试保证双语内容、教学结构和工程质量的一致性。当前项目包含 6 章、32 个知识小节、32 个交互演示、7 个案例阶段和 20 条参考文献。", after=8)
    p = doc.add_paragraph()
    set_paragraph_format(p, 0, 10, 19, False)
    r = p.add_run("关键词：")
    set_run_font(r, "SimSun", 10.5, True)
    r = p.add_run("页面布局；教学网站；交互式可视化；双语内容；响应式设计")
    set_run_font(r, "SimSun", 10.5)

    add_heading(doc, "一 项目概述", 1)
    add_body(doc, "本项目的目标是将“如何做好页面布局”转化为一个能够阅读、操作和验证的教学网站。传统 CSS 教程常以属性语法为主，例如分别介绍 Grid、Flex 或媒体查询的写法，但学习者往往仍然不知道应该在什么场景下选择哪一种规则。本项目因此把重点放在布局决策上：先判断内容是否可读，再建立分组与尺度，随后安排结构和视觉层级，最后处理不同尺寸下的适应性。")
    add_body(doc, "网站以“布局是对注意力的调度”为主线。它不把页面设计理解为单纯的视觉装饰，而是强调正文行长、信息分组、视觉主次、键盘焦点顺序和加载稳定性等可观察结果。用户可以在讲解中理解原理，在代码面板中查看实现方式，再在演示面板中通过滑块、拖拽和模式切换观察参数改变后的实际效果。")

    add_heading(doc, "二 作业要求与完成情况", 1)
    add_body(doc, "本项目针对作业要求进行了逐项落实。网站的内容主题、页面结构和交互方式均围绕页面布局展开，重点不是展示框架本身，而是说明技术实现如何服务于教学目标。")
    add_table(doc,
        ["作业要求", "项目实现", "完成说明"],
        [
            ["教学网站", "课程门户、章节页、案例馆和参考页", "形成完整学习路径，支持从课程入口进入章节、案例和工具页。"],
            ["模仿样例结构", "顶部章节导航、左侧目录、章节翻页和三面板小节", "继承样例的信息组织方式，并使用组件化实现。"],
            ["讲解 关键代码 演示", "每个小节都有 Explain、Key Code、Demo 三个面板", "讲解给出理论与适用边界，代码解释用途和陷阱，演示用于操作验证。"],
            ["动态可交互演示", "32 个 Svelte 与 D3 互动演示", "用户可拖动滑块、切换布局、调整参数或比较阶段，数据和图形会实时变化。"],
            ["中英双语一致", "路由级中英文页面和成对内容模型", "中文与英文内容在同一数据结构中维护，并由脚本检查漏译、引用和演示关联。"],
            ["理论内容而非纯语法", "布局感知、格式塔分组、网格、内在尺寸、可访问性、CLS 等", "每小节至少绑定一条理论文献，避免将课程退化为属性手册。"],
            ["案例式展开", "NeoCampus 新生数据看板 stage0 到 stage6", "同一份 DOM 和数据贯穿六章，逐层增加布局能力，构成完整案例叙事。"],
        ],
        [3.0, 5.2, 7.5],
        9.1,
    )

    add_heading(doc, "三 教学内容与信息架构", 1)
    add_heading(doc, "三点一 三面板教学结构", 2)
    add_body(doc, "每个知识小节采用固定的三面板结构。讲解面板不只描述结论，还说明理论依据、为什么有效以及在什么情况下不宜照搬；关键代码面板为每段代码补充“解决什么问题、关键点在哪里、常见错误是什么”；演示面板使用真实计算而非静态截图。固定结构降低了学习者在不同章节之间切换的认知成本，也使课程内容更容易保持一致。")
    add_body(doc, "三面板由静态 HTML 和局部 Svelte 交互共同实现。无 JavaScript 时三个面板仍可以阅读，启用 JavaScript 后才获得 Tab 切换、键盘操作和演示加载能力。这种渐进增强方式使阅读内容不会依赖复杂的前端状态。")

    add_heading(doc, "三点二 六层课程结构", 2)
    add_body(doc, "六章内容按照依赖关系编排，而不是将六个主题并列排列。内容层解决可读性，间距层建立分组，尺度层统一数值，结构层安排二维和一维布局，视觉层表达业务主次，适应层保证不同尺寸和加载过程中的正确性。前一层为后一层提供条件，减少后续用边框、阴影或大量媒体查询弥补前期结构问题的情况。")
    add_table(doc,
        ["章节", "课程重点", "案例阶段", "代表性交互"],
        [
            ["内容层", "常规流、内在尺寸、正文行长", "stage0 到 stage1", "FlowLab、SizingLab、DensityLab"],
            ["间距层", "格式塔分组、间距标尺、垂直节奏", "stage1 到 stage2", "GestaltLab、SpacingScaleLab、RhythmLab"],
            ["尺度层", "12 列网格、列宽数学、排版音阶", "stage2 到 stage3", "GridBuilder、ColumnMath、TypeScaleLab"],
            ["结构层", "Flex 与 Grid 分工、areas、对齐、DOM 顺序", "stage3 到 stage4", "FlexVsGrid、AreaPainter、TabOrderLab"],
            ["视觉层", "扫视路径、层级、平衡、深度与装饰克制", "stage4 到 stage5", "ScanPathLab、HierarchyLab、BalanceLab"],
            ["适应层", "内容断点、clamp、容器查询、CLS 与性能", "stage5 到 stage6", "BreakpointFinder、ClampPlotter、CLSLab"],
        ],
        [2.1, 5.6, 2.9, 5.1],
        8.9,
    )

    add_heading(doc, "三点三 双语内容设计", 2)
    add_body(doc, "网站不采用在同一页面中保留两套中文和英文 DOM、再通过脚本隐藏其中一套的方式，而是使用路由级国际化。中文页面位于 /zh 路径，英文页面位于 /en 路径。这样可以避免 DOM 体积翻倍，也让每种语言拥有可索引、可直接分享的独立页面。")
    add_body(doc, "在内容层面，所有面向用户的文本以 { zh, en } 的成对结构保存。构建脚本会检查文本是否为空、是否疑似漏译、理论引用是否存在、演示 ID 是否已注册。中英文一致性因此不仅依赖人工检查，也受到类型与自动化脚本的共同约束。")

    add_heading(doc, "四 贯穿案例与量化指标", 1)
    add_body(doc, "NeoCampus 新生数据看板是课程的贯穿案例。看板包含标题、筛选器、摘要、关键指标卡、趋势图、学院分布、表格和操作按钮。stage0 并不是一个故意写错的页面，而是一份具有正确语义结构、仅使用浏览器默认样式的原始数据文档。之后每一章只增加一层新的 CSS 能力，数据和 DOM 保持不变。")
    add_body(doc, "这种案例设计使抽象知识拥有稳定落点。学习者在每一章都能看到页面当前处于什么状态、本章需要增加什么能力、完成后会发生哪些可视变化。案例馆进一步提供七阶段时间轴、可拖动的前后对比、布局决策树和上线前自检清单。")
    add_table(doc,
        ["指标", "含义", "变化结果", "负责层"],
        [
            ["正文行长", "正文单行字符数，舒适区约为 45 到 75ch", "138ch 降至 64ch", "内容层"],
            ["分组比", "组间间距与组内间距的比值", "1.0 提升至 2.0", "间距层"],
            ["取值种类", "页面中可追溯的尺度和取值体系", "3 增至 17", "尺度层"],
            ["可访问性问题", "焦点顺序和过小目标等问题数", "6 降至 0", "结构层和适应层"],
            ["视觉层级分", "视觉重量排序与业务重要度排序的相关性", "0.05 提升至 0.84", "视觉层"],
            ["CLS", "页面加载时的累计布局偏移", "0.31 降至 0.004", "适应层"],
        ],
        [2.2, 5.4, 3.2, 4.9],
        8.9,
    )
    add_body(doc, "这些指标用于教学中的比较和讨论，并不试图将所有审美体验完全量化。例如，视觉层级分通过视觉重量与业务重要度的 Spearman 秩相关衡量主次关系是否一致；行长、CLS、焦点顺序和目标尺寸则分别对应阅读、稳定性和键盘可达性。将部分体验转化为可计算结果，有助于让布局修改具备可解释的依据。")

    add_heading(doc, "五 技术方案与实现特点", 1)
    add_body(doc, "项目采用纯前端静态架构。Astro 负责路由、静态页面生成和内容组织，Svelte 负责具有状态的局部交互。用户首先获得可阅读的静态页面，只有在打开演示或使用对比、时间轴等功能时才加载相应组件。该架构适合以阅读为主、以互动验证为辅的课程网站。")
    add_table(doc,
        ["技术", "主要职责", "对教学网站的价值"],
        [
            ["Astro 7", "路由、静态生成、页面骨架", "生成纯静态页面，降低部署和首屏阅读成本。"],
            ["Svelte 5", "交互孤岛和状态管理", "为 Tab、滑块、案例对比和实验提供局部交互，不将整站做成重型 SPA。"],
            ["TypeScript 与 Zod", "内容模型和构建校验", "约束章节、双语文本、理论引用、代码与演示的完整性。"],
            ["D3.js", "SVG 图表和教学可视化", "展示参数变化后的实时计算结果，适合需要坐标、数据和交互的内容。"],
            ["p5.js", "生成式装饰层", "用于首页和章节页头的非信息性画面，不承担数据表达。"],
            ["Shiki", "代码高亮", "在构建期完成关键代码高亮，减少运行时负担。"],
            ["Vitest", "纯函数单元测试", "验证网格、流体排版、颜色、层级和布局指标算法。"],
        ],
        [2.7, 4.4, 8.6],
        9.0,
    )
    add_heading(doc, "五点一 性能控制", 2)
    add_body(doc, "演示组件统一登记在注册表中，并通过动态 import 按需加载。读者不打开某个演示，就不需要下载该演示的 JavaScript。D3 也不采用整包导入，而是通过子包和统一入口使用所需模块；较少使用的重型模块单独动态加载。这样既保留互动图表，又避免首页承担全部可视化代码。")
    add_body(doc, "p5.js 被限定为装饰层。在用户开启减少动态效果偏好时，p5 不会加载；离开视口或切换到后台标签时动画停止。该策略避免装饰影响内容阅读、可访问性和页面稳定性。")

    add_heading(doc, "六 版式与交互设计", 1)
    add_body(doc, "本项目的页面风格与教学主题保持一致。全站采用 12 列网格、8pt 间距标尺、1.25 排版音阶和 24px 基线节奏。字号、间距、颜色、圆角和阴影集中在设计令牌文件中管理，样式检查脚本会阻止非豁免区域随意使用硬编码字号和间距。这样既保证视觉一致性，也让网站自身成为课程规则的示例。")
    add_body(doc, "网站支持明暗主题。亮色主题以纸面阅读为主要场景，使用阴影表达层次；暗色主题降低对阴影的依赖，改用不同亮度的表面色区分内容层级。用户可在任意页面按 G 显示 12 列网格，按 B 显示 24px 基线，以审查页面是否遵守其教授的布局规则。")
    add_body(doc, "交互设计强调可理解性。章节目录帮助用户定位小节；每章开头的案例入口用原生 dialog 提供当前阶段与目标阶段的对照；三面板支持左右方向键、Home 和 End 键切换；案例馆中的阶段时间轴和对比滑块允许用户以视觉方式观察 CSS 层层叠加的结果。")

    add_heading(doc, "七 可访问性与质量保障", 1)
    add_body(doc, "网站将可访问性作为布局课程的一部分，而不是最后附加的检查项。页面包含跳到主内容链接，三面板遵循 WAI ARIA Tabs 键盘模型，案例弹窗使用浏览器原生 dialog 管理焦点。内容的 DOM 顺序按照业务重要度安排，避免使用 CSS order 制造视觉顺序和键盘焦点顺序不一致的问题。适应层还关注 400% 缩放、最小目标尺寸、容器查询和页面重排。")
    add_body(doc, "工程质量通过多道自动检查维护。Zod 在构建期保证每个小节都具备理论、讲解、代码和演示；双语检查脚本验证中英文内容、文献引用和演示注册关系；设计令牌检查脚本扫描样式文件；Vitest 覆盖布局度量、网格、流体函数和颜色计算。")
    add_table(doc,
        ["验证项目", "当前结果", "说明"],
        [
            ["双语与内容关联检查", "通过", "检查 1214 组双语字符串、20 条文献和 32 个演示。"],
            ["设计令牌检查", "通过", "扫描 29 个样式文件，间距与字号均来自令牌体系。"],
            ["Astro 和 Svelte 检查", "0 errors 0 warnings", "类型检查与组件检查均通过。"],
            ["Vitest 单元测试", "101 项通过", "覆盖视觉重量、行长、CLS、焦点顺序、网格和流体规则等。"],
            ["静态构建", "21 页生成成功", "可输出到 dist 目录并部署到静态托管服务。"],
        ],
        [3.8, 3.6, 8.3],
        9.1,
    )

    add_heading(doc, "八 项目关键亮点", 1)
    add_body(doc, "第一，课程内容通过同一案例形成连续叙事。NeoCampus 看板从原始数据文档出发，六章每次只增加一层能力，学习者能够清楚看到“为什么此时需要这一层”和“加入后解决了什么问题”。这比将 Grid、Flex、间距和响应式各自独立讲解更容易建立整体理解。")
    add_body(doc, "第二，项目把理论、代码和操作验证放在同一个小节中。用户不必只凭阅读记忆概念，而可以直接改变参数，查看图表、布局和指标如何变化。32 个交互演示使教学内容具有实验性质，也符合可视化作业对动态交互的要求。")
    add_body(doc, "第三，网站的工程实现服务于课程主题。设计令牌、网格覆盖、基线覆盖、容器查询、零 order 重排和 CLS 控制并非只在文字中出现，而是直接应用在网站自身。读者可以按键检查实际页面，从而验证课程规则是否真正落地。")
    add_body(doc, "第四，双语和质量保障具有工程化约束。路由级国际化避免双 DOM 冗余，成对的语言结构与自动检查降低漏译风险；Zod、令牌检查、类型检查、单元测试和静态构建共同保障内容质量。这些设计使项目不仅能展示效果，也具备持续维护和扩展的基础。")

    add_heading(doc, "九 总结", 1)
    add_body(doc, "《页面布局之道》完成了一个以案例为中心的页面布局教学网站。项目满足了课程作业关于教学结构、讲解与关键代码、动态演示、中英双语、理论内容和案例展开的要求，并在此基础上加入了可量化指标、内容校验、按需加载、布局审查工具和可访问性设计。最终网站将页面布局从零散 CSS 语法提升为一套能够解释、操作和验证的设计与工程方法。")

    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
